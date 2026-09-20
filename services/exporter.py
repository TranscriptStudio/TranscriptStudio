import json
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer
)
from reportlab.lib.styles import (
    getSampleStyleSheet
)
from reportlab.lib.enums import TA_JUSTIFY

class TranscriptExporter:

    def export_txt(
        self,
        project,
        output_file,
        include_speakers=True,
        include_timestamps=False
    ):

        with open(
            output_file,
            "w",
            encoding="utf-8"
        ) as f:

            for block in project.blocks:

                speaker = (
                    project.get_speaker_by_id(
                        block.speaker_id
                    )
                )

                text = (
                    project.get_block_text(
                        block
                    )
                )

                line = ""

                if include_timestamps:

                    start_word = (
                        project.words[
                            block.start_word
                        ]
                    )

                    t = int(
                        start_word.start or 0
                    )

                    h = t // 3600
                    m = (
                        t % 3600
                    ) // 60
                    s = t % 60

                    line += (
                        f"[{h:02}:{m:02}:{s:02}] "
                    )

                if include_speakers:

                    line += (
                        f"[{speaker.name}]"
                    )

                if line:

                    f.write(line + "\n")

                f.write(text)

                f.write("\n \n")

    def export_json(
        self,
        project,
        output_file,
        include_speakers=True,
        include_timestamps=False
    ):

        data = []

        for block in project.blocks:

            entry = {}

            speaker = (
                project.get_speaker_by_id(
                    block.speaker_id
                )
            )

            if include_speakers:

                entry["speaker"] = (
                    speaker.name
                )

            if include_timestamps:

                entry["start"] = (
                    project.words[
                        block.start_word
                    ].start
                )

            entry["text"] = (
                project.get_block_text(
                    block
                )
            )

            data.append(entry)

        with open(
            output_file,
            "w",
            encoding="utf-8"
        ) as f:

            json.dump(
                data,
                f,
                ensure_ascii=False,
                indent=2
            )

    def export_docx(
        self,
        project,
        output_file,
        include_speakers=True,
        include_timestamps=False
    ):

        document = Document()

        for block in project.blocks:

            speaker = project.get_speaker_by_id(
                block.speaker_id
            )

            text = project.get_block_text(
                block
            )

            header = ""

            if include_timestamps:

                start_word = project.words[
                    block.start_word
                ]

                t = int(start_word.start or 0)

                h = t // 3600
                m = (t % 3600) // 60
                s = t % 60

                header += (
                    f"[{h:02}:{m:02}:{s:02}] "
                )

            if include_speakers:

                header += (
                    f"[{speaker.name}]"
                )

            if header:

                p = document.add_paragraph()

                run = p.add_run(header)

                run.bold = True

            p = document.add_paragraph(text)
            p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY

            document.add_paragraph("")

        document.save(output_file)

    def export_pdf(
        self,
        project,
        output_file,
        include_speakers=True,
        include_timestamps=False
    ):

        doc = SimpleDocTemplate(
            output_file
        )

        styles = getSampleStyleSheet()
        body_style = styles["BodyText"]
        body_style.alignment = TA_JUSTIFY

        story = []

        for block in project.blocks:

            speaker = project.get_speaker_by_id(
                block.speaker_id
            )

            text = project.get_block_text(
                block
            )

            header = ""

            if include_timestamps:

                start_word = project.words[
                    block.start_word
                ]

                t = int(start_word.start or 0)

                h = t // 3600
                m = (t % 3600) // 60
                s = t % 60

                header += (
                    f"[{h:02}:{m:02}:{s:02}] "
                )

            if include_speakers:

                header += (
                    f"[{speaker.name}]"
                )

            if header:

                story.append(
                    Paragraph(
                        f"<b>{header}</b>",
                        styles["Heading4"]
                    )
                )

            story.append(
                Paragraph(
                    text,
                    body_style
                )
            )

            story.append(
                Spacer(
                    1,
                    12
                )
            )

        doc.build(story)






