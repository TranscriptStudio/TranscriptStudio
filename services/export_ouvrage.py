# services/export_ouvrage.py
"""
Export de l'ouvrage (P18.3 — feuille de route §7).

Assemblage des textes validés dans l'ordre de
l'étude :

- par réunion (ordre de lecture) : résumé validé en
  tête de section (si souhaité), puis les
  reformulations validées dans l'**ordre d'ouvrage**
  de la réunion ;
- DOCX via python-docx : styles sobres, guillemets
  “...”, références normalisées, en-tête de réunion
  optionnel, **pas de gras dans le corps** ;
- export TXT simple en complément.

Le flux de fin réutilise celui de l'export v1
(P13) : message de fin avec ouverture du fichier ou
du dossier (côté UI).
"""

import re
from pathlib import Path


def _derniere_version_validee(resultat):
    """
    Dernière version validée d'un résultat
    (réformulation ou résumé), ou None.
    """

    if not resultat or not resultat.versions:
        return None

    for version in reversed(resultat.versions):

        if (
            version.validation
            and version.validation.statut == "valide"
        ):
            return version

    return None


def normaliser_guillemets(texte: str) -> str:
    """
    Convention de l'ouvrage (§7) : guillemets
    anglais “...” au lieu des guillemets français.
    """

    if not texte:
        return texte

    texte = re.sub(
        r"«\s*",
        "“",
        texte
    )

    return re.sub(
        r"\s*»",
        "”",
        texte
    )


def assembler_ouvrage(
    etude_ouverte,
    inclure_resumes: bool = True,
    en_tetes_reunions: bool = True
) -> dict:
    """
    Assemble l'ouvrage à partir des textes validés.

    Retour :
    ``{"titre", "reunions": [{"titre", "resume",
    "blocs": [textes...]}]}`` — seules les réunions
    comportant au moins un texte validé figurent dans
    l'assemblage.
    """

    from services.decoupage import (
        ordre_ouvrage_effectif,
    )

    titre_etude = (
        etude_ouverte.etude.titre
        or etude_ouverte.etude.theme
        or "Ouvrage"
    )

    sections = []

    for nom_reunion in (
        etude_ouverte.etude.ordre_reunions
    ):

        reunion = etude_ouverte.reunions.get(
            nom_reunion
        )

        titre_reunion = (
            (reunion.titre_1 if reunion else "")
            or nom_reunion
        )

        #
        # Résumé validé (tête de section, si inclus).
        #
        resume = None

        if inclure_resumes:

            resultat_resume = (
                etude_ouverte.resumes.get(nom_reunion)
            )

            version = _derniere_version_validee(
                resultat_resume
            )

            if version is not None:
                resume = version.texte

        #
        # Reformulations validées dans l'ordre
        # d'ouvrage de la réunion.
        #
        blocs = []

        try:

            ordre = ordre_ouvrage_effectif(
                etude_ouverte,
                nom_reunion
            )

        except Exception:
            ordre = []

        for fid in ordre:

            resultat = (
                etude_ouverte.reformulations.get(fid)
            )

            version = _derniere_version_validee(
                resultat
            )

            if version is not None:
                blocs.append(version.texte)

        if not blocs and resume is None:
            continue

        sections.append({
            "titre": titre_reunion,
            "resume": resume,
            "blocs": blocs,
        })

    return {
        "titre": titre_etude,
        "reunions": sections,
        "en_tetes_reunions": en_tetes_reunions,
    }


def _paragraphes(texte: str) -> list:
    """
    Paragraphes d'un texte validé : un seul retour à
    la ligne par saut de paragraphe (contrat v2).
    """

    if not texte:
        return []

    return [
        normaliser_guillemets(paragraphe)
        for paragraphe in texte.split("\n")
        if paragraphe.strip()
    ]


def exporter_docx(ouvrage: dict, chemin) -> Path:
    """
    Export DOCX de l'ouvrage : styles sobres, pas de
    gras dans le corps (les titres sont des styles
    de structure).
    """

    from docx import Document
    from docx.shared import Pt

    document = Document()

    document.add_heading(ouvrage["titre"], level=0)

    for section in ouvrage["reunions"]:

        if ouvrage.get("en_tetes_reunions"):

            document.add_heading(
                section["titre"],
                level=1
            )

        for paragraphe in _paragraphes(
            section["resume"]
        ):
            document.add_paragraph(paragraphe)

        for texte in section["blocs"]:

            for paragraphe in _paragraphes(texte):
                document.add_paragraph(paragraphe)

            document.add_paragraph("")

    chemin = Path(chemin)

    document.save(str(chemin))

    return chemin


def exporter_txt(ouvrage: dict, chemin) -> Path:
    """
    Export TXT simple de l'ouvrage.
    """

    lignes = [ouvrage["titre"], ""]

    for section in ouvrage["reunions"]:

        if ouvrage.get("en_tetes_reunions"):
            lignes.append(section["titre"])
            lignes.append("=" * len(section["titre"]))
            lignes.append("")

        for paragraphe in _paragraphes(
            section["resume"]
        ):
            lignes.append(paragraphe)

            lignes.append("")

        for texte in section["blocs"]:

            lignes.extend(_paragraphes(texte))

            lignes.append("")

    chemin = Path(chemin)

    chemin.write_text(
        "\n".join(lignes).strip() + "\n",
        encoding="utf-8"
    )

    return chemin
