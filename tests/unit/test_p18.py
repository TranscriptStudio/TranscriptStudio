"""
Tests P18 — JND locale, validation des références
et citations, export de l'ouvrage (feuille de
route §5.2, §5.4, §6, §7).

- P18.1 : chargeur OSIS XML (avec espace de noms),
  lookup par référence, recherche, parsing /
  normalisation des références françaises ;
- P18.2 : contrôle qualité — références complétées
  (suggestion + extrait JND), références perdues
  (alerte discrète), citations exactes ou
  s'écartant du texte JND, alignement À LA DEMANDE ;
- P18.3 : assemblage de l'ouvrage (ordre d'étude,
  résumés en tête de section, ordre d'ouvrage) et
  export DOCX (pas de gras dans le corps,
  guillemets “...”) + TXT.
"""

import os
from pathlib import Path

os.environ.setdefault(
    "QT_QPA_PLATFORM",
    "offscreen"
)

import pytest

from services.controle_qualite import (
    aligner_citations,
    controles_bloc,
    controle_references,
)
from services.jnd import (
    BibleAbsente,
    charger_bible,
    extraire_references,
    normaliser_reference,
)

OSIS_TEST = """<?xml version="1.0" encoding="UTF-8"?>
<osis xmlns="http://www.crosswire.org/OSIS/namespace">
  <osisText osisIDWork="JND">
    <div type="book" osisID="Gen">
      <chapter osisID="Gen 1">
        <verse osisID="Gen 1.1">Au commencement, Dieu créa les cieux et la terre.</verse>
        <verse osisID="Gen 1.2">Et la terre était désolation et vide.</verse>
      </chapter>
      <chapter osisID="Gen 2">
        <verse osisID="Gen 2.4">Voici les origines des cieux et de la terre.</verse>
        <verse osisID="Gen 2.5">Aucun arbuste des champs n'était encore sur la terre.</verse>
        <verse osisID="Gen 2.6">Mais une vapeur montait de la terre<note>note de l'éditeur</note> et arrosait toute la surface du sol.</verse>
      </chapter>
    </div>
    <div type="book" osisID="Luke">
      <chapter osisID="Luke 15">
        <verse osisID="Luke 15.11">Un homme avait deux fils.</verse>
        <verse osisID="Luke 15.12">Le plus jeune dit à son père : Donne-moi la part de bien qui me revient.</verse>
      </chapter>
    </div>
  </osisText>
</osis>
"""


@pytest.fixture
def bible(tmp_path):

    chemin = tmp_path / "JND.osis.xml"
    chemin.write_text(OSIS_TEST, encoding="utf-8")

    return charger_bible(chemin)


#
# Format milestone des modules réels (CrossWire :
# Darby / DarbyR) : le texte VIT ENTRE
# <verse sID/> et <verse eID/> — un lecteur naïf
# n'y verrait que des éléments vides.
#

OSIS_MILESTONE = """<?xml version="1.0" encoding="UTF-8"?>
<osis xmlns="http://www.bibletechnologies.net/2003/OSIS/namespace">
  <osisText osisIDWork="FreJND">
    <div type="book" osisID="Luke">
      <chapter osisID="Luke 15">
        <verse sID="Luke.15.11" osisID="Luke.15.11" n="11"/>Et il dit\u00a0: <w lemma="strong:G0444">Un</w> homme avait deux fils\u202f;<note>note de l'éditeur à exclure</note> et le plus jeune partit.
        <verse eID="Luke.15.11"/>
        <verse sID="Luke.15.12" osisID="Luke.15.12" n="12"/>Le plus jeune dit à son père\u00a0: Donne-moi la part.
        <verse eID="Luke.15.12"/>
      </chapter>
    </div>
  </osisText>
</osis>
"""


@pytest.fixture
def bible_milestone(tmp_path):

    chemin = tmp_path / "Darby.osis.xml"
    chemin.write_text(OSIS_MILESTONE, encoding="utf-8")

    return charger_bible(chemin)


class TestJNDMilestones:

    def test_texte_entre_les_marqueurs(self, bible_milestone):

        ref = extraire_references("Luc 15.11")[0]

        texte = bible_milestone.texte_reference(ref)

        # Texte réel (tail du sID + <w> + tails), la
        # note exclue, espaces insécables nettoyées.
        assert texte == (
            "Et il dit : Un homme avait deux fils ; "
            "et le plus jeune partit."
        )

    def test_plage_et_recherche(self, bible_milestone):

        ref = extraire_references("Luc 15.11-12")[0]

        texte = bible_milestone.texte_reference(ref)

        assert "deux fils" in texte
        assert "Donne-moi la part" in texte

        resultats = bible_milestone.rechercher(
            "Donne-moi"
        )

        assert resultats[0]["reference"] == (
            "Luc 15.12"
        )

    def test_fichier_choisi_par_priorite(self, tmp_path):

        #
        # Un JND.osis.xml explicite prime sur les
        # modules CrossWire ; sinon la version par
        # défaut de P18.5 (DarbyR) avant la Darby
        # classique (repli implicite uniquement —
        # le contrôle qualité passe par la version
        # retenue par l'étude).
        #
        from services import jnd as service_jnd

        (tmp_path / "DarbyR.osis.xml").write_text(
            OSIS_MILESTONE,
            encoding="utf-8"
        )

        chemin = service_jnd.chemin_bible()

        # Sans JND.osis.xml dans le dossier des
        # ressources réelles, le fichier local
        # n'est PAS pris : chemin_bible regarde
        # resources/bible/. On teste donc la
        # priorité relative des fichiers réels.
        noms = service_jnd.NOMS_FICHIERS_BIBLE

        assert noms.index("JND.osis.xml") < (
            noms.index("DarbyR.osis.xml")
        ) < noms.index("Darby.osis.xml")


@pytest.fixture
def etude_validee(tmp_path):
    """
    Étude d'une réunion : blocs 1 et 2 validés (pas
    le 3), résumé validé, ordre d'ouvrage [2, 1, 3],
    plus une 2e réunion sans rien validé.
    """

    from services.etude_io import (
        ajouter_version_reformulation,
        ajouter_version_resume,
        valider_version_reformulation,
        valider_version_resume,
    )
    from services.decoupage import (
        definir_ordre_ouvrage,
    )
    from tests.unit.test_reformulation_ui import (
        construire_etude_test,
        construire_projet_test,
    )
    from services.etude_io import (
        ajouter_reunion,
        sauvegarder_etude,
    )

    _projet, _projet_v1, etude, nom = (
        construire_etude_test(tmp_path)
    )

    for fid in (1, 2):
        ajouter_version_reformulation(
            etude,
            fid,
            f"Texte validé du bloc {fid}.",
            type_version="passe",
        )

        version = (
            etude.reformulations[fid]
            .versions[-1]
        )

        valider_version_reformulation(
            etude,
            fid,
            numero=version.numero,
            validateur="Test"
        )

    ajouter_version_resume(
        etude,
        nom,
        "Résumé validé de la réunion.",
        source="reformulations"
    )

    valider_version_resume(
        etude,
        nom,
        numero=1,
        validateur="Test"
    )

    definir_ordre_ouvrage(etude, nom, [2, 1, 3])

    #
    # Deuxième réunion (aucun texte validé) : elle
    # ne doit pas figurer dans l'ouvrage.
    #
    (tmp_path / "deuxieme").mkdir()

    construire_projet_test(tmp_path / "deuxieme")

    projet2 = (
        tmp_path / "deuxieme" / "2026-10-13_Luc1b"
    )

    projet2_renomme = (
        tmp_path / "deuxieme" / "2026-10-20_Luc2"
    )

    projet2.rename(projet2_renomme)

    ajouter_reunion(
        etude,
        projet2_renomme,
        "Deuxième chapitre"
    )

    sauvegarder_etude(etude)

    return etude, nom


#
# ==================================================
# P18.1 — ressource JND.
# ==================================================
#

class TestJND:

    def test_absente(self, tmp_path):

        with pytest.raises(BibleAbsente):

            charger_bible(
                tmp_path / "inconnu.osis.xml"
            )

    def test_lookup_ponctuel_et_plage(self, bible):

        ref = extraire_references("Genèse 1.1")[0]

        assert "Au commencement" in (
            bible.texte_reference(ref)
        )

        ref = extraire_references("Genèse 2.4-6")[0]

        texte = bible.texte_reference(ref)

        assert "Voici les origines" in texte
        assert "vapeur" in texte

        # Les <note> OSIS sont exclues du texte.
        assert "note de l'éditeur" not in texte

    def test_lookup_chapitre_entier(self, bible):

        ref = extraire_references("Genèse 1")[0]

        texte = bible.texte_reference(ref)

        assert "Au commencement" in texte
        assert "désolation" in texte

    def test_reference_introuvable(self, bible):

        ref = extraire_references("Osée 3.1")[0]

        assert bible.texte_reference(ref) == ""

    def test_parsing_variantes(self):

        cas = [
            ("genèse 2 verset 4 à 6", "Genèse 2.4-6"),
            ("Genèse 2 versets 4-6", "Genèse 2.4-6"),
            ("Luc 15.11-32", "Luc 15.11-32"),
            ("Luc 15:11", "Luc 15.11"),
            ("Jn 3:16", "Jean 3.16"),
            ("Ps 23", "Psaumes 23"),
            ("1 Samuel 3", "1 Samuel 3"),
            ("Esaie 53", "Ésaïe 53"),
            ("Luc 15", "Luc 15"),
        ]

        for texte, attendu in cas:

            references = extraire_references(
                f"… comme dit en {texte} …"
            )

            assert references, texte

            assert str(references[0]) == attendu

    def test_pas_de_faux_positifs(self):

        references = extraire_references(
            "Il a fait un bon job et quelques actes "
            "de gentillesse."
        )

        assert references == []

    def test_normaliser_reference(self):

        assert normaliser_reference(
            "genèse 2 verset 4 à 6"
        ) == "Genèse 2.4-6"

        assert normaliser_reference(
            "aucune référence ici"
        ) is None

    def test_rechercher(self, bible):

        resultats = bible.rechercher("commencement")

        assert resultats

        assert resultats[0]["reference"] == (
            "Genèse 1.1"
        )

        assert "Au commencement" in (
            resultats[0]["texte"]
        )


#
# ==================================================
# P18.2 — contrôle qualité des références.
# ==================================================
#

class TestControleReferences:

    def test_reference_completee_avec_extrait(self, bible):

        controle = controle_references(
            "Le fils prodigue, nous l'avons vu en "
            "Luc 15.",
            "Le fils prodigue (cf. Luc 15.11-12) "
            "parle de grâce.",
            bible=bible
        )

        assert len(controle["ajoutees"]) == 1

        ajoutee = controle["ajoutees"][0]

        assert ajoutee["reference"] == (
            "Luc 15.11-12"
        )

        assert "deux fils" in ajoutee["extrait_jnd"]

        assert "complétées" in controle["message"]

    def test_completion_versets_suggestion(self, bible):

        controle = controle_references(
            "Nous avons lu Luc 15 ensemble.",
            "Nous avons lu le chapitre du père et "
            "des deux fils (Luc 15.11-12).",
            bible=bible
        )

        # Complétion de versets : suggestion de
        # vérification (les numéros ont été
        # ajoutés) — pas une perte.
        assert len(controle["ajoutees"]) == 1
        assert controle["perdues"] == []

        assert "deux fils" in controle["ajoutees"][
            0
        ]["extrait_jnd"]

    def test_meme_niveau_rien_a_signaler(self, bible):

        controle = controle_references(
            "Nous avons lu Luc 15 ensemble.",
            "Nous avons relu Luc 15 aujourd'hui.",
            bible=bible
        )

        assert controle["ajoutees"] == []
        assert controle["perdues"] == []

    def test_reference_perdue(self, bible):

        controle = controle_references(
            "Comme en Jean 3.16, Dieu a aimé.",
            "Dieu a aimé le monde.",
            bible=bible
        )

        assert len(controle["perdues"]) == 1

        assert controle["perdues"][0][
            "reference"
        ] == "Jean 3.16"

        assert "absentes du résultat" in (
            controle["message"]
        )

    def test_sans_bible_parsing_local(self):

        controle = controle_references(
            "Texte source.",
            "Texte avec une référence à Marc 2.5.",
            bible=None
        )

        assert len(controle["ajoutees"]) == 1

        assert "extrait_jnd" not in (
            controle["ajoutees"][0]
        )

        assert "README" in controle["message"]

    def test_controles_bloc_complet(self, bible):

        controles = controles_bloc(
            "Source : nous avons lu Luc 15.",
            "Résultat : Luc 15.11-12 et Jean 3.16 "
            "parlent de grâce.",
            {},
            bible=bible
        )

        assert "message_references" in controles

        assert controles["nb_alertes"] >= 1


#
# ==================================================
# P18.2 — citations et alignement (§5.4).
# ==================================================
#

class TestCitations:

    def test_citation_exacte(self, bible):

        texte = (
            "Jésus dit : « Un homme avait deux "
            "fils. » — c'est Luc 15.11."
        )

        from services.jnd import (
            extraire_references,
        )
        from services.controle_qualite import (
            controle_citations,
        )

        citations = controle_citations(
            texte,
            references=extraire_references(texte),
            bible=bible
        )

        assert citations

        assert citations[0]["statut"] == "exacte"

    def test_citation_paraphrase_ecart(self, bible):

        texte = (
            "Jésus dit : « Il y avait un père avec "
            "deux garçons. » — c'est Luc 15.11."
        )

        from services.jnd import (
            extraire_references,
        )
        from services.controle_qualite import (
            controle_citations,
        )

        citations = controle_citations(
            texte,
            references=extraire_references(texte),
            bible=bible
        )

        assert citations

        assert citations[0]["statut"] == "ecart"

        assert "deux fils" in citations[0][
            "texte_jnd"
        ]

    def test_alignement_a_la_demande(self, bible):

        texte = (
            "Jésus dit : « Il y avait un père avec "
            "deux garçons. » — c'est Luc 15.11. "
            "Et il ajoute : « Un homme avait deux "
            "fils. » en Luc 15.11."
        )

        aligne = aligner_citations(texte, bible)

        # La paraphrase est alignée sur la JND...
        assert "« Un homme avait deux fils. »" in (
            aligne
        )

        # ...une seule fois : le second passage
        # (déjà exact) n'est pas dupliqué.
        assert aligne.count("deux fils") == (
            texte.count("deux fils") + 1
        )

    def test_alignement_aucun_ecart(self, bible):

        texte = (
            "Jésus dit : « Un homme avait deux "
            "fils. » — c'est Luc 15.11."
        )

        assert aligner_citations(
            texte,
            bible
        ) == texte


#
# ==================================================
# P18.3 — export de l'ouvrage.
# ==================================================
#

class TestExportOuvrage:

    def test_assemblage_ordre_ouvrage(
        self,
        etude_validee
    ):

        from services.export_ouvrage import (
            assembler_ouvrage,
        )

        etude, nom = etude_validee

        ouvrage = assembler_ouvrage(etude)

        assert ouvrage["titre"] == (
            "Évangile de Luc"
        )

        # La réunion sans texte validé est exclue.
        assert len(ouvrage["reunions"]) == 1

        section = ouvrage["reunions"][0]

        # Résumé en tête de section.
        assert section["resume"] == (
            "Résumé validé de la réunion."
        )

        # Ordre d'ouvrage [2, 1, 3] : seuls les
        # blocs 1 et 2 sont validés — dans cet ordre.
        assert section["blocs"] == [
            "Texte validé du bloc 2.",
            "Texte validé du bloc 1.",
        ]

    def test_assemblage_options(self, etude_validee):

        from services.export_ouvrage import (
            assembler_ouvrage,
        )

        etude, nom = etude_validee

        ouvrage = assembler_ouvrage(
            etude,
            inclure_resumes=False,
            en_tetes_reunions=False
        )

        section = ouvrage["reunions"][0]

        assert section["resume"] is None
        assert ouvrage["en_tetes_reunions"] is False

    def test_export_docx(self, etude_validee, tmp_path):

        from docx import Document

        from services.export_ouvrage import (
            assembler_ouvrage,
            exporter_docx,
        )

        etude, nom = etude_validee

        ouvrage = assembler_ouvrage(etude)

        chemin = tmp_path / "ouvrage.docx"

        exporter_docx(ouvrage, chemin)

        assert chemin.is_file()

        document = Document(str(chemin))

        textes = [
            paragraphe.text
            for paragraphe in document.paragraphs
        ]

        # Ordre : résumé, puis bloc 2, puis bloc 1.
        ordre = [
            index
            for index, texte in enumerate(textes)
            if texte.startswith(
                ("Résumé", "Texte validé")
            )
        ]

        assert textes[ordre[0]].startswith("Résumé")

        assert textes[ordre[1]] == (
            "Texte validé du bloc 2."
        )

        assert textes[ordre[2]] == (
            "Texte validé du bloc 1."
        )

        # Pas de gras dans le corps.
        for paragraphe in document.paragraphs:

            if not paragraphe.text.startswith(
                ("Résumé", "Texte validé")
            ):
                continue

            for run in paragraphe.runs:
                assert not run.bold

    def test_export_txt(self, etude_validee, tmp_path):

        from services.export_ouvrage import (
            assembler_ouvrage,
            exporter_txt,
        )

        etude, nom = etude_validee

        ouvrage = assembler_ouvrage(etude)

        chemin = tmp_path / "ouvrage.txt"

        exporter_txt(ouvrage, chemin)

        contenu = chemin.read_text(encoding="utf-8")

        assert "Évangile de Luc" in contenu
        assert "Résumé validé" in contenu

        index_2 = contenu.index(
            "Texte validé du bloc 2."
        )
        index_1 = contenu.index(
            "Texte validé du bloc 1."
        )

        assert index_2 < index_1

    def test_guillemets_ouvrage(self, etude_validee, tmp_path):

        from services.etude_io import (
            ajouter_version_reformulation,
            valider_version_reformulation,
        )
        from services.export_ouvrage import (
            assembler_ouvrage,
            exporter_docx,
        )
        from docx import Document

        etude, nom = etude_validee

        ajouter_version_reformulation(
            etude,
            3,
            "Il dit : « grâce sur grâce. »",
            type_version="edition_manuelle",
        )

        valider_version_reformulation(
            etude,
            3,
            numero=1,
            validateur="Test"
        )

        ouvrage = assembler_ouvrage(etude)

        chemin = tmp_path / "guillemets.docx"

        exporter_docx(ouvrage, chemin)

        document = Document(str(chemin))

        textes = [
            paragraphe.text
            for paragraphe in document.paragraphs
        ]

        assert any(
            "“grâce sur grâce.”" in texte
            for texte in textes
        )

        assert not any(
            "«" in texte for texte in textes
        )
